"""Build the Excel version of the model, for colleagues who do not code.

Every result in the workbook is a live Excel formula. Change a number on the
"Inputs" sheet (for example the Ideal TP) and the whole model recalculates.

Sheets:
  Read Me       : what the model does, a glossary, and the data disclaimer
  Inputs        : every setting (blue cells can be edited)
  District Data : one row per district (inputs from the data pipeline)
  Model         : all calculations, one row per district
  Summary       : totals by expansion wave, region and rollout year
  Macro Results : filled in by the VBA macro (vba/FootprintModel.bas)
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from . import config

FONT = "Arial"
BLUE_INPUT = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN_LINK = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
WHITE_BOLD = Font(name=FONT, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, color="595959", size=9)
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
KEY_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")

DISCLAIMER = (
    "All sales, retailer and store data in this workbook is SIMULATED. It is not real company data "
    "and exists only to explain the method. District boundaries (geoBoundaries) and population "
    "(Pakistan Population Census 2023) are public."
)

FIRST = 5  # first data row on District Data and Model


def _name(wb, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _header(ws, row: int, headers: list[str], widths: dict | None = None) -> None:
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = WHITE_BOLD
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 45
    for i, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 14)


def _readme(wb) -> None:
    ws = wb.active
    ws.title = "Read Me"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "Retail Footprint Expansion Model"
    ws["A1"].font = TITLE
    ws["A2"] = DISCLAIMER
    ws["A2"].font = Font(name=FONT, bold=True, color="C00000")
    ws.merge_cells("A2:B2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 45
    rows = [
        ("What this model answers", "How many new freezers should the business invest in, in each district, and do enough real shops exist to take them?"),
        ("Who it is for", "Anyone in Sales, Finance, Marketing or Operations. No coding needed: change the blue cells on the Inputs sheet and everything recalculates."),
        ("How to use it", "1. Review or change the settings on the Inputs sheet.  2. Read the results on the Model and Summary sheets.  3. Optional: run the macro RunFootprintModel (see vba/FootprintModel.bas) to recalculate outlet recommendations step by step on the Macro Results sheet."),
        ("Color code", "Blue text on yellow = input you can change. Black text = formula. Green text = link to another sheet."),
        ("", ""),
        ("Glossary", ""),
        ("PCC", "Per Capita Consumption: liters of ice cream sold per target person per year (Volume / Target population)."),
        ("PPO", "Population Per Outlet: target people for every freezer outlet (Target population / Outlets). High PPO means underserved."),
        ("TP", "Throughput: liters sold per freezer per week (Volume / Outlets / Weeks). One outlet means one freezer."),
        ("NI", "New Induction: a retailer that received a freezer during the year. NI TP is the throughput of those new retailers."),
        ("LSM 5+", "Living Standards Measure 5 and above: the share of people who can afford to buy ice cream regularly (the target population)."),
        ("GOLY", "Growth Over Last Year."),
        ("CAGR", "Compound Annual Growth Rate (2022 to 2024)."),
        ("Ideal TP / Ideal PPO", "Benchmarks. Freezers are added until TP falls to the Ideal TP or PPO falls to the Ideal PPO, whichever comes first."),
        ("First fill", "Liters needed to stock a brand new freezer the first time."),
        ("Census universe", "Stores from a third party retail census that sell the main snacking competitors, are in a target channel, have high turnover and are not already customers."),
        ("Conversion rate", "Share of census universe stores expected to agree to take a freezer."),
        ("Expansion wave", "Order of expansion from the strategy: 1 Strongholds (citadel and distributor based districts), 2 their Neighbouring districts, 3 Population High districts, 4 the Rest."),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=a).font = BOLD
        c = ws.cell(row=i, column=2, value=b)
        c.font = BLACK
        c.alignment = WRAP


def _inputs(wb) -> None:
    ws = wb.create_sheet("Inputs")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "Inputs (edit the blue cells)"
    ws["A1"].font = TITLE
    ws["A2"] = DISCLAIMER
    ws["A2"].font = NOTE
    items = [
        ("IdealTP", "Ideal TP (liters per freezer per week)", config.IDEAL_TP, "Stop adding freezers when throughput falls to this level.", "0.0"),
        ("IdealPPO", "Ideal PPO (target people per outlet)", config.IDEAL_PPO, "Industry benchmark.", "#,##0"),
        ("Weeks", "Weeks per year", config.WEEKS_PER_YEAR, "Volumes are full year volumes.", "0"),
        ("FirstFill", "First fill per new freezer (liters)", config.FIRST_FILL_LITERS, "Stock needed the first time a freezer is placed.", "0"),
        ("TPBasis", "Throughput assumed for each new outlet", "NI TP", "NI TP = throughput of last year's new retailers (original method). District TP = average of all retailers.", "@"),
        ("ConvRate", "Census conversion rate", config.CENSUS_CONVERSION_RATE, "Share of eligible census stores expected to accept a freezer.", "0%"),
        ("TopN", "Number of districts in the shortlist", config.TOP_N_DISTRICTS, "Top districts by priority.", "0"),
        ("Budget1", "Budget scenario 1 (freezers funded)", config.BUDGET_SCENARIOS[0], "Shared across shortlisted districts by contribution.", "#,##0"),
        ("Budget2", "Budget scenario 2 (freezers funded)", config.BUDGET_SCENARIOS[1], "Shared across shortlisted districts by contribution.", "#,##0"),
    ]
    years = list(config.PHASING)
    for i, y in enumerate(years, start=1):
        items.append((f"Share{y}", f"Rollout share in {y}", config.PHASING[y], "Original plan split: 300, 300 and 250 out of 850.", "0.0%"))
    for i, (cut, label) in enumerate(config.PCC_LAYERS, start=1):
        items.append((f"PCC_L{i}", f"PCC threshold for {label}", cut, "Liters per target person per year (at or above).", "0.00"))
    for i, (cut, label) in enumerate(config.PPO_LAYERS, start=1):
        items.append((f"PPO_L{i}", f"PPO threshold for {label}", cut, "Target people per outlet (at or above).", "#,##0"))

    ws["A4"], ws["B4"], ws["C4"] = "Setting", "Value", "Note"
    for c in ("A4", "B4", "C4"):
        ws[c].font = WHITE_BOLD
        ws[c].fill = HEADER_FILL
    for r, (name, label, value, note, fmt) in enumerate(items, start=5):
        ws.cell(row=r, column=1, value=label).font = BLACK
        v = ws.cell(row=r, column=2, value=value)
        v.font = BLUE_INPUT
        v.fill = INPUT_FILL
        v.number_format = fmt
        ws.cell(row=r, column=3, value=note).font = NOTE
        _name(wb, name, f"Inputs!$B${r}")
        if name == "TPBasis":
            dv = DataValidation(type="list", formula1='"NI TP,District TP"', allow_blank=False)
            ws.add_data_validation(dv)
            dv.add(v)

    # Expansion wave lookup table
    start = 5 + len(items) + 2
    ws.cell(row=start - 1, column=1, value="Expansion wave by nature of district").font = BOLD
    for r, (nature, wave) in enumerate(config.WAVE_ORDER.items(), start=start):
        ws.cell(row=r, column=1, value=nature).font = BLACK
        c = ws.cell(row=r, column=2, value=wave)
        c.font = BLUE_INPUT
        c.fill = INPUT_FILL
    end = start + len(config.WAVE_ORDER) - 1
    _name(wb, "WaveNames", f"Inputs!$A${start}:$A${end}")
    _name(wb, "WaveNumbers", f"Inputs!$B${start}:$B${end}")


DATA_COLS = [
    ("District", "district", None),
    ("Province", "province", None),
    ("Sales region", "region", None),
    ("Nature of district", "nature_of_district", None),
    ("Population 2024 (census based)", "population_2024", "#,##0"),
    ("LSM 5+ share", "lsm5_share", "0.0%"),
    ("Retailers 2022", "retailers_2022", "#,##0"),
    ("Retailers 2023", "retailers_2023", "#,##0"),
    ("Retailers 2024", "retailers_2024", "#,##0"),
    ("Volume 2022 (liters)", "volume_2022", "#,##0"),
    ("Volume 2023 (liters)", "volume_2023", "#,##0"),
    ("Volume 2024 (liters)", "volume_2024", "#,##0"),
    ("NI retailers 2024", "ni_retailers_2024", "#,##0"),
    ("NI volume 2024 (liters)", "ni_volume_2024", "#,##0"),
    ("Census universe stores (eligible, not yet customers)", "census_universe", "#,##0"),
]


def _district_data(wb, df) -> int:
    ws = wb.create_sheet("District Data")
    ws["A1"] = "District Data (inputs produced by the Python pipeline)"
    ws["A1"].font = TITLE
    ws["A2"] = DISCLAIMER
    ws["A2"].font = NOTE
    _header(ws, FIRST - 1, [c[0] for c in DATA_COLS], {1: 26, 4: 16, 15: 20})
    for r, row in enumerate(df.itertuples(index=False), start=FIRST):
        for c, (_, field, fmt) in enumerate(DATA_COLS, start=1):
            cell = ws.cell(row=r, column=c, value=getattr(row, field))
            cell.font = BLUE_INPUT
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    return FIRST + len(df) - 1


# Model sheet: (header, formula template, number format). {r} is the row,
# {D} prefixes a District Data reference, {last} is the last data row.
def _model_columns():
    D = "'District Data'!"
    rng = lambda col: f"${col}${FIRST}:${col}${{last}}"  # noqa: E731
    return [
        ("District", f"={D}A{{r}}", None),
        ("Sales region", f"={D}C{{r}}", None),
        ("Nature of district", f"={D}D{{r}}", None),
        ("Target population 2024 (LSM 5+)", f"=ROUND({D}E{{r}}*{D}F{{r}},0)", "#,##0"),
        ("Outlets 2024", f"={D}I{{r}}", "#,##0"),
        ("Volume 2024 (liters)", f"={D}L{{r}}", "#,##0"),
        ("TP 2024 (liters per freezer per week)", "=IFERROR(F{r}/(E{r}*Weeks),0)", "0.0"),
        ("NI TP 2024", f"=IFERROR({D}N{{r}}/({D}M{{r}}*Weeks),0)", "0.0"),
        ("Region NI TP (used when a district had no NI)",
         f"=IFERROR(SUMIFS({D}{rng('N')},{D}{rng('C')},B{{r}})/(SUMIFS({D}{rng('M')},{D}{rng('C')},B{{r}})*Weeks),0)", "0.0"),
        ("TP assumed per new outlet", '=IF(TPBasis="NI TP",IF(H{r}>0,H{r},I{r}),G{r})', "0.0"),
        ("PCC 2024 (liters per target person)", "=IFERROR(F{r}/D{r},0)", "0.000"),
        ("PPO 2024", "=IFERROR(D{r}/E{r},0)", "#,##0"),
        ("Volume GOLY", f"=IFERROR({D}L{{r}}/{D}K{{r}}-1,0)", "0.0%"),
        ("Volume CAGR 2022 to 2024", f"=IFERROR(({D}L{{r}}/{D}J{{r}})^(1/2)-1,0)", "0.0%"),
        ("Outlets GOLY", f"=IFERROR({D}I{{r}}/{D}H{{r}}-1,0)", "0.0%"),
        ("PCC layer", '=IF(K{r}>=PCC_L1,"Layer 01",IF(K{r}>=PCC_L2,"Layer 02",IF(K{r}>=PCC_L3,"Layer 03",IF(K{r}>=PCC_L4,"Layer 04","Layer 05"))))', None),
        ("PPO layer", '=IF(L{r}>=PPO_L1,"Layer 01",IF(L{r}>=PPO_L2,"Layer 02",IF(L{r}>=PPO_L3,"Layer 03",IF(L{r}>=PPO_L4,"Layer 04","Layer 05"))))', None),
        ("Outlets to add before PPO reaches ideal", "=D{r}/IdealPPO-E{r}", "#,##0.0"),
        ("TP gap per new outlet (liters per year)", "=IdealTP*Weeks-(Weeks*J{r}+FirstFill)", "#,##0.0"),
        ("Outlets to add before TP reaches ideal", "=IF(S{r}>0,(F{r}-IdealTP*Weeks*E{r})/S{r},1E+12)", "#,##0.0"),
        ("RECOMMENDED NEW OUTLETS (Phase 1 demand)",
         "=IF(E{r}<=0,0,IF(G{r}<=IdealTP,0,IF(L{r}<=IdealPPO,0,INT(MIN(T{r},R{r})+0.000000001))))", "#,##0"),
        ("Limiting factor",
         '=IF(E{r}<=0,"No current outlets",IF(G{r}<=IdealTP,"TP already at or below ideal",IF(L{r}<=IdealPPO,"PPO already at or below ideal",IF(T{r}<=R{r},"Stopped by TP reaching ideal","Stopped by PPO reaching ideal"))))', None),
        ("New total volume (liters)", "=F{r}+U{r}*(Weeks*J{r}+FirstFill)", "#,##0"),
        ("New total outlets", "=E{r}+U{r}", "#,##0"),
        ("New TP", "=IFERROR(W{r}/(X{r}*Weeks),0)", "0.0"),
        ("New PPO", "=IFERROR(D{r}/X{r},0)", "#,##0"),
        ("Quadrant (PCC against PPO)",
         '=IF(AND(K{r}>=NationalPCC,L{r}>IdealPPO),"Q1: High PCC, High PPO",IF(K{r}>=NationalPCC,"Q2: High PCC, Low PPO",IF(L{r}>IdealPPO,"Q3: Low PCC, High PPO","Q4: Low PCC, Low PPO")))', None),
        ("Phase", '=IF(K{r}>=NationalPCC,IF(L{r}>IdealPPO,"Phase 1","Phase 2"),"Phase 3")', None),
        ("Quadrant (TP against PPO)",
         '=IF(AND(G{r}>IdealTP,L{r}>IdealPPO),"High TP, High PPO",IF(G{r}>IdealTP,"High TP, Low PPO",IF(L{r}>IdealPPO,"Low TP, High PPO","Low TP, Low PPO")))', None),
        ("Expansion wave", "=INDEX(WaveNumbers,MATCH(C{r},WaveNames,0))", "0"),
        ("Sort score (helper)", '=IF(U{r}>0,AD{r}*10000000+VALUE(RIGHT(AB{r},1))*1000000+(1000000-U{r})+ROW()/1000000,"")', "0.000000"),
        ("Priority rank", f'=IF(U{{r}}>0,COUNTIFS({rng("AE")},"<"&AE{{r}})+1,"")', "0"),
        ("Shortlisted", '=IF(AF{r}="","No",IF(AF{r}<=TopN,"Yes","No"))', None),
        ("Usable census supply", f"=INT({D}O{{r}}*ConvRate)", "#,##0"),
        ("FINAL NEW OUTLETS (smaller of demand and supply)", "=MIN(U{r},AH{r})", "#,##0"),
        ("Constraint",
         '=IF(U{r}=0,"No expansion",IF(AH{r}<U{r},"Supply constrained (not enough shops)","Demand constrained (enough shops)"))', None),
        ("Final, shortlisted only", '=IF(AG{r}="Yes",AI{r},0)', "#,##0"),
    ]


def _budget_and_phasing_columns(start_col: int):
    """Largest remainder splits, written as plain formulas (no macros needed)."""
    L = get_column_letter
    cols = []
    c = start_col
    for b in (1, 2):
        exact, base, frac, alloc = L(c), L(c + 1), L(c + 2), L(c + 3)
        cols += [
            (f"Budget {b}: exact share (helper)",
             f'=IF(AG{{r}}="Yes",IF(Budget{b}>=ShortlistFinal,AK{{r}},AK{{r}}/ShortlistFinal*Budget{b}),0)', "0.000"),
            (f"Budget {b}: whole part (helper)", f"=INT({exact}{{r}})", "#,##0"),
            (f"Budget {b}: remainder (helper)",
             f'=IF(AND(AG{{r}}="Yes",Budget{b}<ShortlistFinal),{exact}{{r}}-{base}{{r}},-1)', "0.000"),
            (f"BUDGET SCENARIO {b} FREEZERS",
             f'=IF(AG{{r}}<>"Yes",0,IF(Budget{b}>=ShortlistFinal,AK{{r}},{base}{{r}}+IF(COUNTIFS(${frac}${FIRST}:${frac}${{last}},">"&{frac}{{r}})'
             f'+COUNTIFS(${frac}${FIRST}:${frac}${{last}},{frac}{{r}},$AF${FIRST}:$AF${{last}},"<"&AF{{r}})'
             f"<Budget{b}-SUM(${base}${FIRST}:${base}${{last}}),1,0)))", "#,##0"),
        ]
        c += 4
    years = list(config.PHASING)
    e = [L(c + i) for i in range(3)]
    bse = [L(c + 3 + i) for i in range(3)]
    f = [L(c + 6 + i) for i in range(3)]
    short = L(c + 9)
    for i, y in enumerate(years):
        cols.append((f"{y} exact (helper)", f"=AK{{r}}*Share{y}", "0.000"))
    for i, y in enumerate(years):
        cols.append((f"{y} whole (helper)", f"=INT({e[i]}{{r}})", "0"))
    for i, y in enumerate(years):
        cols.append((f"{y} remainder (helper)", f"={e[i]}{{r}}-{bse[i]}{{r}}", "0.000"))
    cols.append(("Left to place (helper)", f"=AK{{r}}-{bse[0]}{{r}}-{bse[1]}{{r}}-{bse[2]}{{r}}", "0"))
    ahead = [
        f"({f[1]}{{r}}>{f[0]}{{r}})+({f[2]}{{r}}>{f[0]}{{r}})",
        f"({f[0]}{{r}}>{f[1]}{{r}})+({f[2]}{{r}}>{f[1]}{{r}})+({f[0]}{{r}}={f[1]}{{r}})",
        f"({f[0]}{{r}}>{f[2]}{{r}})+({f[1]}{{r}}>{f[2]}{{r}})+({f[0]}{{r}}={f[2]}{{r}})+({f[1]}{{r}}={f[2]}{{r}})",
    ]
    for i, y in enumerate(years):
        cols.append((f"ROLLOUT {y}", f"={bse[i]}{{r}}+IF({ahead[i]}<{short}{{r}},1,0)", "#,##0"))
    return cols


def _model(wb, last: int) -> None:
    ws = wb.create_sheet("Model")
    ws["A1"] = "Model (all formulas)"
    ws["A1"].font = TITLE
    cols = _model_columns()
    cols += _budget_and_phasing_columns(len(cols) + 1)
    _header(ws, FIRST - 1, [c[0] for c in cols], {1: 26, 3: 16, 21: 16, 22: 30, 27: 24, 29: 20, 36: 34})

    ws["C2"] = "National PCC"
    ws["D2"] = f"=SUM(F{FIRST}:F{last})/SUM(D{FIRST}:D{last})"
    ws["D2"].number_format = "0.000"
    ws["E2"] = "Shortlist final total"
    ws["F2"] = f"=SUM(AK{FIRST}:AK{last})"
    ws["F2"].number_format = "#,##0"
    for c in ("C2", "E2"):
        ws[c].font = BOLD
    _name(wb, "NationalPCC", "Model!$D$2")
    _name(wb, "ShortlistFinal", "Model!$F$2")

    key_cols = {i for i, c in enumerate(cols, start=1) if c[0].isupper() or c[0].startswith(("RECOMMENDED", "FINAL", "BUDGET", "ROLLOUT"))}
    for r in range(FIRST, last + 1):
        for i, (_, tmpl, fmt) in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i, value=tmpl.format(r=r, last=last))
            cell.font = GREEN_LINK if i <= 6 else BLACK
            if fmt:
                cell.number_format = fmt
            if i in key_cols:
                cell.fill = KEY_FILL
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    ws.conditional_formatting.add(
        f"AG{FIRST}:AG{last}", CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="C6E0B4"))
    )


def _summary(wb, last: int) -> None:
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18
    ws["A1"] = "Summary"
    ws["A1"].font = TITLE
    ws["A2"] = DISCLAIMER
    ws["A2"].font = NOTE
    M = "Model!"
    rng = lambda col: f"{M}${col}${FIRST}:${col}${last}"  # noqa: E731

    ws["A4"] = "Headline"
    ws["A4"].font = BOLD
    head = [
        ("Districts in the model", f"=COUNTA({rng('A')})"),
        ("Current outlets (freezers)", f"=SUM({rng('E')})"),
        ("Districts with demand for new outlets", f'=COUNTIF({rng("U")},">0")'),
        ("Shortlisted districts", f'=COUNTIF({rng("AG")},"Yes")'),
        ("Phase 1 demand, all districts (new outlets)", f"=SUM({rng('U')})"),
        ("Phase 1 demand, shortlisted districts", f'=SUMIFS({rng("U")},{rng("AG")},"Yes")'),
        ("Phase 2 usable supply, shortlisted districts", f'=SUMIFS({rng("AH")},{rng("AG")},"Yes")'),
        ("FINAL new outlets, shortlisted districts", "=ShortlistFinal"),
        ("Supply constrained shortlisted districts", f'=COUNTIFS({rng("AG")},"Yes",{rng("AJ")},"Supply*")'),
    ]
    for i, (label, f) in enumerate(head, start=5):
        ws.cell(row=i, column=1, value=label).font = BLACK
        c = ws.cell(row=i, column=2, value=f)
        c.number_format = "#,##0"
        c.font = BLACK

    r0 = 5 + len(head) + 1
    ws.cell(row=r0, column=1, value="By expansion wave (shortlisted)").font = BOLD
    for j, h in enumerate(["Districts", "Demand", "Usable supply", "Final"], start=2):
        ws.cell(row=r0, column=j, value=h).font = BOLD
    for i, (nature, _) in enumerate(config.WAVE_ORDER.items(), start=r0 + 1):
        ws.cell(row=i, column=1, value=nature).font = BLACK
        crit = f'{rng("C")},$A{i},{rng("AG")},"Yes"'
        ws.cell(row=i, column=2, value=f"=COUNTIFS({crit})")
        ws.cell(row=i, column=3, value=f"=SUMIFS({rng('U')},{crit})")
        ws.cell(row=i, column=4, value=f"=SUMIFS({rng('AH')},{crit})")
        ws.cell(row=i, column=5, value=f"=SUMIFS({rng('AI')},{crit})")
        for j in range(2, 6):
            ws.cell(row=i, column=j).number_format = "#,##0"

    r1 = r0 + len(config.WAVE_ORDER) + 2
    ws.cell(row=r1, column=1, value="By sales region (shortlisted)").font = BOLD
    for j, h in enumerate(["Districts", "Final", f"Budget {config.BUDGET_SCENARIOS[0]:,}", f"Budget {config.BUDGET_SCENARIOS[1]:,}"], start=2):
        ws.cell(row=r1, column=j, value=h).font = BOLD
    budget_cols = ["AO", "AS"]
    for i, region in enumerate(sorted(_REGIONS), start=r1 + 1):
        ws.cell(row=i, column=1, value=region).font = BLACK
        crit = f'{rng("B")},$A{i},{rng("AG")},"Yes"'
        ws.cell(row=i, column=2, value=f"=COUNTIFS({crit})")
        ws.cell(row=i, column=3, value=f"=SUMIFS({rng('AI')},{crit})")
        ws.cell(row=i, column=4, value=f"=SUMIFS({rng(budget_cols[0])},{crit})")
        ws.cell(row=i, column=5, value=f"=SUMIFS({rng(budget_cols[1])},{crit})")
        for j in range(2, 6):
            ws.cell(row=i, column=j).number_format = "#,##0"

    r2 = r1 + len(_REGIONS) + 2
    ws.cell(row=r2, column=1, value="Rollout by year (final plan)").font = BOLD
    rollout_cols = ["BD", "BE", "BF"]
    for i, (y, col) in enumerate(zip(config.PHASING, rollout_cols), start=r2 + 1):
        ws.cell(row=i, column=1, value=str(y)).font = BLACK
        c = ws.cell(row=i, column=2, value=f"=SUM({rng(col)})")
        c.number_format = "#,##0"
    tot = r2 + len(config.PHASING) + 1
    ws.cell(row=tot, column=1, value="Total").font = BOLD
    ws.cell(row=tot, column=2, value=f"=SUM(B{r2 + 1}:B{tot - 1})").number_format = "#,##0"

    r3 = tot + 2
    ws.cell(row=r3, column=1, value="Budget scenarios").font = BOLD
    for i, (b, col) in enumerate(zip((1, 2), budget_cols), start=r3 + 1):
        ws.cell(row=i, column=1, value=f"Budget scenario {b}").font = BLACK
        ws.cell(row=i, column=2, value=f"=Budget{b}").number_format = "#,##0"
        ws.cell(row=i, column=3, value=f"=SUM({rng(col)})").number_format = "#,##0"
    ws.cell(row=r3, column=2, value="Funded").font = BOLD
    ws.cell(row=r3, column=3, value="Allocated").font = BOLD


def _macro_sheet(wb) -> None:
    ws = wb.create_sheet("Macro Results")
    ws["A1"] = "Macro Results"
    ws["A1"].font = TITLE
    ws["A2"] = "Run the macro RunFootprintModel (vba/FootprintModel.bas) to fill this sheet. It adds outlets one at a time, replacing manual Goal Seek, and checks its answer against the Model sheet."
    ws["A2"].font = NOTE
    _header(ws, 4, ["District", "Recommended new outlets (macro)", "Limiting factor (macro)", "Model sheet formula", "Match?"],
            {1: 26, 2: 18, 3: 32, 4: 18, 5: 10})


_REGIONS: list[str] = []


def build_workbook(district_model, path) -> None:
    """Write the Excel model. district_model is the Python pipeline's district table."""
    df = district_model.sort_values("district").reset_index(drop=True).copy()
    _REGIONS[:] = sorted(df["region"].unique())
    wb = Workbook()
    _readme(wb)
    _inputs(wb)
    last = _district_data(wb, df)
    _model(wb, last)
    _summary(wb, last)
    _macro_sheet(wb)
    # Make sure every cell uses the same professional font
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != FONT:
                    f = cell.font
                    cell.font = Font(name=FONT, bold=f.bold, italic=f.italic, size=f.size, color=f.color)
    wb["Summary"].sheet_properties.tabColor = "1F3864"
    wb.calculation.fullCalcOnLoad = True  # Excel recalculates every formula when the file opens
    wb.save(path)
